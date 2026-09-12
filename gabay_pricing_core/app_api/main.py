from __future__ import annotations

import hashlib
import json
import os
from contextlib import contextmanager
from pathlib import Path
from typing import Any

import openpyxl

try:
    # Local dev convenience only: loads MONDAY_API_TOKEN and friends from the
    # repo-root .env (gitignored, never committed) without overriding any
    # variable the process already has set. A missing python-dotenv or .env
    # file is not fatal -- the app still runs, just without those optional
    # integrations (see monday_integration.MondayConfigError).
    from dotenv import load_dotenv

    load_dotenv(Path(__file__).resolve().parents[2] / ".env", override=False)
except ImportError:
    pass
from fastapi import FastAPI, File, HTTPException, Query, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError

from pricing_core import normalize_inventory_rows

from . import pricing_service
from .database import Base, build_database
from .db_models import (
    InventoryVersionRow,
    MarketSnapshotRow,
    PricingSessionRow,
    ProjectRow,
    ProjectSaleRow,
    ProjectStateSnapshotRow,
    ProjectStateUnitRow,
    ScenarioRow,
    SourceRunRow,
    StrategyProfileRow,
    UnitRow,
)
from .demo_snapshot import RequiredSourceMissingError, build_demo_market_snapshot
from .inventory_preview import build_inventory_preview
from .migrations import apply_schema_migrations
from .market_context_registry import MARKET_CONTEXTS
from .market_context_workspace import build_market_context_workspace_payload
from .monday_integration import MondayApiError, MondayConfigError, create_pricing_approval_item
from .petah_tikva_workspace import build_petah_tikva_scenario_payload, build_petah_tikva_workspace_payload
from .project_launcher import resolve_project_start
from .schemas import (
    FamilyDecisionUpdate,
    InventoryImport,
    InventoryVersionRead,
    LifecycleEventCreate,
    MondayPricingApprovalRequest,
    PetahTikvaScenarioRequest,
    PricingSessionCreate,
    PricingSessionRead,
    ProjectCreate,
    ProjectRead,
    ProjectSaleCreate,
    ProjectSaleRead,
    ProjectStartRequest,
    ScenarioCreate,
    ScenarioRead,
    StateSnapshotCreate,
    UnitRead,
)

DEMO_PROJECT_KEY = "assignment_demo_project"


def create_app(database_url: str | None = None) -> FastAPI:
    database_url = database_url or os.getenv("DATABASE_URL", "sqlite:///./gabay_pricing.db")
    engine, SessionLocal = build_database(database_url)
    Base.metadata.create_all(engine)
    apply_schema_migrations(engine)

    app = FastAPI(title="Project Pricing Workspace API", version="0.1.0")
    app.add_middleware(
        CORSMiddleware,
        allow_origins=os.getenv("CORS_ALLOW_ORIGINS", "http://localhost:3000").split(","),
        allow_credentials=True,
        allow_methods=["*"],
        allow_headers=["*"],
    )

    @contextmanager
    def session_scope():
        session = SessionLocal()
        try:
            yield session
        finally:
            session.close()

    @app.get("/health")
    def health():
        return {"status": "ok", "pricing_core": "market-range-v1", "api": "v1"}

    @app.post("/api/v1/projects", response_model=ProjectRead, status_code=status.HTTP_201_CREATED)
    def create_project(payload: ProjectCreate):
        with session_scope() as db:
            row = ProjectRow(**payload.model_dump())
            db.add(row)
            db.commit()
            db.refresh(row)
            return row

    @app.get("/api/v1/projects/{project_id}", response_model=ProjectRead)
    def get_project(project_id: str):
        with session_scope() as db:
            row = db.get(ProjectRow, project_id)
            if row is None:
                raise HTTPException(status_code=404, detail="project_not_found")
            return row

    @app.post("/api/v1/demo-project", response_model=ProjectRead)
    def get_or_create_demo_project():
        """Idempotent bootstrap: always keyed by demo_key, never by display name, so
        repeated demo runs reuse the same project instead of creating duplicates."""
        with session_scope() as db:
            row = db.scalar(select(ProjectRow).where(ProjectRow.demo_key == DEMO_PROJECT_KEY))
            if row is None:
                row = ProjectRow(
                    demo_key=DEMO_PROJECT_KEY,
                    name="Assignment Demo Project",
                    city="אשקלון",
                    neighborhood="עיר היין",
                    location_source="candidate_demo_neighborhood_centroid_from_current_listings",
                )
                db.add(row)
                db.commit()
                db.refresh(row)
            return row

    @app.post(
        "/api/v1/projects/{project_id}/inventory/import",
        response_model=InventoryVersionRead,
        status_code=status.HTTP_201_CREATED,
    )
    def import_inventory(project_id: str, payload: InventoryImport):
        with session_scope() as db:
            return _import_matrix(db, project_id, payload.source_filename, payload.matrix)

    @app.post(
        "/api/v1/projects/{project_id}/inventory/import-file",
        response_model=InventoryVersionRead,
        status_code=status.HTTP_201_CREATED,
    )
    async def import_inventory_file(project_id: str, file: UploadFile = File(...)):
        content = await file.read()
        try:
            import io
            workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(status_code=422, detail=f"could_not_read_xlsx: {exc}") from exc
        worksheet = workbook[workbook.sheetnames[0]]
        matrix = [list(row) for row in worksheet.iter_rows(values_only=True)]
        with session_scope() as db:
            return _import_matrix(db, project_id, file.filename or "inventory.xlsx", matrix)

    @app.get("/api/v1/inventory/{inventory_version_id}", response_model=InventoryVersionRead)
    def get_inventory(inventory_version_id: str):
        with session_scope() as db:
            row = db.get(InventoryVersionRow, inventory_version_id)
            if row is None:
                raise HTTPException(status_code=404, detail="inventory_version_not_found")
            return _inventory_read(row, reused_existing=False)

    # --- Generic project-start launcher --------------------------------------
    # City-agnostic entry point. Today it only resolves the one registered
    # Petah Tikva / חפץ חיים 25 frozen snapshot (see app_api.project_launcher);
    # unsupported (city, address) pairs are reported explicitly rather than
    # silently falling back to Petah Tikva. No live collection happens here.

    @app.post("/api/v1/inventory/preview")
    async def inventory_preview(file: UploadFile = File(...)):
        """Stateless read of an uploaded workbook: parses/normalizes it (same
        pricing_core.normalize_inventory_rows the persisted import path uses)
        and returns unit counts/routes only. Writes nothing to the database --
        this never creates a ProjectRow or InventoryVersionRow, so it cannot
        attach the legacy Ashkelon demo project (see get_or_create_demo_project)
        the way the old upload flow on the homepage does."""
        content = await file.read()
        try:
            import io
            workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(status_code=422, detail=f"could_not_read_xlsx: {exc}") from exc
        worksheet = workbook[workbook.sheetnames[0]]
        matrix = [list(row) for row in worksheet.iter_rows(values_only=True)]
        try:
            return build_inventory_preview(matrix)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc

    @app.post("/api/v1/demo/project-start")
    def project_start(payload: ProjectStartRequest):
        try:
            return resolve_project_start(
                payload.city, payload.address, inventory_fingerprint=payload.inventory_fingerprint
            )
        except FileNotFoundError as exc:
            raise HTTPException(status_code=503, detail=f"snapshot_artifact_missing: {exc}") from exc

    # --- Petah Tikva standard-unit demo workspace (frozen artifacts only) --------

    @app.get("/api/v1/demo/petah-tikva/workspace")
    def petah_tikva_workspace():
        """Single-payload read for the Petah Tikva standard-unit workspace screen.

        Reads only the frozen data/frozen/petah_tikva_*_v1.json artifacts already
        produced by the dedicated run_petah_tikva_*_v1.py scripts -- no live
        source fetch, no DB write, so it keeps working even if external services
        are unavailable. Does not recompute any pricing_core methodology.
        """
        try:
            return build_petah_tikva_workspace_payload()
        except FileNotFoundError as exc:
            raise HTTPException(status_code=503, detail=f"petah_tikva_frozen_artifact_missing: {exc}") from exc

    @app.post("/api/v1/demo/petah-tikva/scenario")
    def petah_tikva_scenario(payload: PetahTikvaScenarioRequest):
        """On-demand company-strategy scenario for the 32 standard units, priced
        via the same pricing_core engine as the frozen baseline. Purely
        computed in memory and returned -- never written to data/frozen/*.json,
        so the frozen baseline is never overwritten by this endpoint."""
        try:
            return build_petah_tikva_scenario_payload(payload.range_position_pct, payload.name)
        except FileNotFoundError as exc:
            raise HTTPException(status_code=503, detail=f"petah_tikva_frozen_artifact_missing: {exc}") from exc

    # --- Multi-market-context demo workspace (Petah Tikva + 3 frozen contexts) ---

    @app.get("/api/v1/demo/market-contexts")
    def market_contexts():
        """Lists the market contexts the "הקשר שוק" selector can switch
        between -- summary only, no workspace body (see the per-slug route
        below for the full payload)."""
        return [
            {"slug": c.slug, "display_name": c.display_name, "city": c.city, "submarket": c.submarket}
            for c in MARKET_CONTEXTS.values()
        ]

    @app.get("/api/v1/demo/market-contexts/{slug}/workspace")
    def market_context_workspace(slug: str):
        """Single-payload read for any of the four market contexts, in the
        same schema build_petah_tikva_workspace_payload already produces
        (plus one additive `market_context` key). Petah Tikva's own slug
        delegates to that unmodified function directly -- see
        market_context_workspace.build_market_context_workspace_payload."""
        context = MARKET_CONTEXTS.get(slug)
        if context is None:
            raise HTTPException(status_code=404, detail=f"unknown_market_context: {slug}")
        try:
            return build_market_context_workspace_payload(context)
        except FileNotFoundError as exc:
            raise HTTPException(status_code=503, detail=f"market_context_frozen_artifact_missing: {exc}") from exc

    # --- Monday.com pricing-approval integration ---------------------------------
    # One outbound demo integration: create a real item on the existing
    # "אישור מחירון – פרויקט פתח תקווה" board so approval continues there. No
    # inbound sync/webhooks (see app_api.monday_integration's module docstring).

    @app.post("/api/v1/integrations/monday/pricing-approval")
    def monday_pricing_approval(payload: MondayPricingApprovalRequest):
        try:
            return create_pricing_approval_item(payload)
        except MondayConfigError as exc:
            raise HTTPException(status_code=503, detail="monday_not_configured") from exc
        except MondayApiError as exc:
            raise HTTPException(status_code=502, detail="monday_request_failed") from exc

    # --- market snapshots --------------------------------------------------------

    @app.post("/api/v1/projects/{project_id}/demo-market-snapshot", status_code=status.HTTP_201_CREATED)
    def create_demo_market_snapshot(project_id: str):
        with session_scope() as db:
            project = db.get(ProjectRow, project_id)
            if project is None:
                raise HTTPException(status_code=404, detail="project_not_found")
            try:
                snapshot = build_demo_market_snapshot(db, project)
                db.commit()
            except RequiredSourceMissingError as exc:
                db.rollback()
                raise HTTPException(status_code=422, detail=f"required_source_missing: {exc}") from exc
            return _snapshot_read(db, snapshot)

    @app.get("/api/v1/market-snapshots/{snapshot_id}")
    def get_market_snapshot(snapshot_id: str):
        with session_scope() as db:
            snapshot = db.get(MarketSnapshotRow, snapshot_id)
            if snapshot is None:
                raise HTTPException(status_code=404, detail="market_snapshot_not_found")
            return _snapshot_read(db, snapshot)

    # --- pricing sessions ---------------------------------------------------------

    @app.post("/api/v1/pricing-sessions", response_model=PricingSessionRead, status_code=status.HTTP_201_CREATED)
    def create_pricing_session(payload: PricingSessionCreate):
        with session_scope() as db:
            if db.get(ProjectRow, payload.project_id) is None:
                raise HTTPException(status_code=404, detail="project_not_found")
            if db.get(InventoryVersionRow, payload.inventory_version_id) is None:
                raise HTTPException(status_code=404, detail="inventory_version_not_found")
            if db.get(MarketSnapshotRow, payload.market_snapshot_id) is None:
                raise HTTPException(status_code=404, detail="market_snapshot_not_found")

            state_snapshot_id = payload.project_state_snapshot_id
            if state_snapshot_id is not None:
                if db.get(ProjectStateSnapshotRow, state_snapshot_id) is None:
                    raise HTTPException(status_code=404, detail="project_state_snapshot_not_found")
            else:
                snapshot = pricing_service.create_state_snapshot(db, payload.project_id, payload.inventory_version_id)
                state_snapshot_id = snapshot.id

            row = PricingSessionRow(
                project_id=payload.project_id,
                inventory_version_id=payload.inventory_version_id,
                market_snapshot_id=payload.market_snapshot_id,
                project_state_snapshot_id=state_snapshot_id,
            )
            db.add(row)
            db.commit()
            db.refresh(row)
            return row

    @app.get("/api/v1/pricing-sessions/compare")
    def compare_pricing_sessions(
        session_a: str = Query(...),
        scenario_a: str = Query(...),
        session_b: str = Query(...),
        scenario_b: str = Query(...),
    ):
        # Registered before the /{session_id} route below on purpose: Starlette
        # matches routes in registration order, and {session_id} would otherwise
        # greedily capture the literal "compare" path segment.
        with session_scope() as db:
            sa = db.get(PricingSessionRow, session_a)
            sb = db.get(PricingSessionRow, session_b)
            ca = db.get(ScenarioRow, scenario_a)
            cb = db.get(ScenarioRow, scenario_b)
            if sa is None or sb is None or ca is None or cb is None:
                raise HTTPException(status_code=404, detail="session_or_scenario_not_found")
            return pricing_service.compare_pricing_sessions(db, sa, ca, sb, cb)

    @app.get("/api/v1/pricing-sessions/{session_id}", response_model=PricingSessionRead)
    def get_pricing_session(session_id: str):
        with session_scope() as db:
            row = db.get(PricingSessionRow, session_id)
            if row is None:
                raise HTTPException(status_code=404, detail="pricing_session_not_found")
            return row

    @app.get("/api/v1/pricing-sessions/{session_id}/scenarios", response_model=list[ScenarioRead])
    def list_scenarios(session_id: str):
        with session_scope() as db:
            if db.get(PricingSessionRow, session_id) is None:
                raise HTTPException(status_code=404, detail="pricing_session_not_found")
            return db.scalars(
                select(ScenarioRow).where(ScenarioRow.pricing_session_id == session_id).order_by(ScenarioRow.created_at)
            ).all()

    @app.post(
        "/api/v1/pricing-sessions/{session_id}/scenarios",
        response_model=ScenarioRead,
        status_code=status.HTTP_201_CREATED,
    )
    def create_scenario(session_id: str, payload: ScenarioCreate):
        with session_scope() as db:
            session_row = db.get(PricingSessionRow, session_id)
            if session_row is None:
                raise HTTPException(status_code=404, detail="pricing_session_not_found")

            parent = None
            if payload.parent_scenario_id:
                parent = db.get(ScenarioRow, payload.parent_scenario_id)
                if parent is None or parent.pricing_session_id != session_id:
                    raise HTTPException(status_code=404, detail="parent_scenario_not_found")

            scenario = ScenarioRow(
                pricing_session_id=session_id,
                parent_scenario_id=payload.parent_scenario_id,
                name=payload.name,
                created_by=payload.created_by,
            )
            db.add(scenario)
            db.flush()

            if parent is not None:
                for decision in db.scalars(
                    select(StrategyProfileRow).where(StrategyProfileRow.scenario_id == parent.id)
                ).all():
                    db.add(StrategyProfileRow(
                        scenario_id=scenario.id,
                        family_key=decision.family_key,
                        name=decision.name,
                        basis=decision.basis,
                        range_position_pct=decision.range_position_pct,
                        competitor_reference_ils=decision.competitor_reference_ils,
                        competitor_reference_source_id=decision.competitor_reference_source_id,
                        competitor_reference_name=decision.competitor_reference_name,
                        competitor_delta_ils=decision.competitor_delta_ils,
                        negotiation_buffer_ils=decision.negotiation_buffer_ils,
                        floor_rule_json=decision.floor_rule_json,
                        minimum_price_ils=decision.minimum_price_ils,
                        maximum_price_ils=decision.maximum_price_ils,
                        internal_sale_plus_amount_ils=decision.internal_sale_plus_amount_ils,
                        minimum_not_below_last_realized_sale=decision.minimum_not_below_last_realized_sale,
                        rationale=decision.rationale,
                        source=decision.source,
                        note=decision.note,
                    ))

            db.commit()
            db.refresh(scenario)
            return scenario

    # --- family decisions + reprice ------------------------------------------------

    @app.put("/api/v1/scenarios/{scenario_id}/family-decisions/{family_id}")
    def upsert_family_decision(scenario_id: str, family_id: str, payload: FamilyDecisionUpdate):
        with session_scope() as db:
            scenario = db.get(ScenarioRow, scenario_id)
            if scenario is None:
                raise HTTPException(status_code=404, detail="scenario_not_found")

            row = db.scalar(
                select(StrategyProfileRow).where(
                    StrategyProfileRow.scenario_id == scenario_id,
                    StrategyProfileRow.family_key == family_id,
                )
            )
            if row is None:
                row = StrategyProfileRow(scenario_id=scenario_id, family_key=family_id)
                db.add(row)

            row.name = payload.name
            row.basis = payload.basis
            row.rationale = payload.rationale
            row.range_position_pct = payload.range_position_pct
            row.competitor_reference_ils = payload.competitor_reference_ils
            row.competitor_reference_source_id = payload.competitor_reference_source_id
            row.competitor_reference_name = payload.competitor_reference_name
            row.competitor_delta_ils = payload.competitor_delta_ils
            row.negotiation_buffer_ils = payload.negotiation_buffer_ils
            row.floor_rule_json = (
                json.dumps(payload.floor_rule.model_dump(), ensure_ascii=False) if payload.floor_rule else None
            )
            row.minimum_price_ils = payload.minimum_price_ils
            row.maximum_price_ils = payload.maximum_price_ils
            row.internal_sale_plus_amount_ils = payload.internal_sale_plus_amount_ils
            row.minimum_not_below_last_realized_sale = payload.minimum_not_below_last_realized_sale
            row.source = payload.source
            row.note = payload.note

            try:
                db.commit()
            except IntegrityError as exc:
                db.rollback()
                raise HTTPException(status_code=409, detail="family_decision_conflict") from exc

            return {"scenario_id": scenario_id, "family_id": family_id, "status": "saved"}

    @app.post("/api/v1/scenarios/{scenario_id}/reprice")
    def reprice(scenario_id: str):
        with session_scope() as db:
            scenario = db.get(ScenarioRow, scenario_id)
            if scenario is None:
                raise HTTPException(status_code=404, detail="scenario_not_found")
            try:
                result = pricing_service.reprice_scenario(db, scenario)
                db.commit()
            except ValueError as exc:
                db.rollback()
                raise HTTPException(status_code=422, detail=str(exc)) from exc
            return result.public_dict()

    @app.get("/api/v1/scenarios/{scenario_id}/price-list")
    def price_list(scenario_id: str):
        with session_scope() as db:
            scenario = db.get(ScenarioRow, scenario_id)
            if scenario is None:
                raise HTTPException(status_code=404, detail="scenario_not_found")
            return pricing_service.get_price_list(db, scenario)

    @app.get("/api/v1/scenarios/{scenario_id}/families")
    def families(scenario_id: str):
        with session_scope() as db:
            scenario = db.get(ScenarioRow, scenario_id)
            if scenario is None:
                raise HTTPException(status_code=404, detail="scenario_not_found")
            return pricing_service.get_families(db, scenario)

    @app.get("/api/v1/scenarios/{scenario_id}/units/{unit_number}/evidence")
    def unit_evidence(scenario_id: str, unit_number: str):
        with session_scope() as db:
            scenario = db.get(ScenarioRow, scenario_id)
            if scenario is None:
                raise HTTPException(status_code=404, detail="scenario_not_found")
            return pricing_service.unit_evidence_payload(db, scenario, unit_number)

    @app.get("/api/v1/scenarios/{scenario_id}/impact")
    def impact(scenario_id: str, against: str = Query(...)):
        with session_scope() as db:
            scenario = db.get(ScenarioRow, scenario_id)
            baseline = db.get(ScenarioRow, against)
            if scenario is None or baseline is None:
                raise HTTPException(status_code=404, detail="scenario_not_found")
            return pricing_service.impact_payload(db, scenario, baseline)

    # --- project lifecycle / own sales / state snapshots --------------------------

    @app.post(
        "/api/v1/projects/{project_id}/sales",
        response_model=ProjectSaleRead,
        status_code=status.HTTP_201_CREATED,
    )
    def create_sale(project_id: str, payload: ProjectSaleCreate):
        with session_scope() as db:
            if db.get(ProjectRow, project_id) is None:
                raise HTTPException(status_code=404, detail="project_not_found")
            try:
                sale = pricing_service.record_sale(db, project_id, payload.model_dump())
                db.commit()
            except pricing_service.DuplicateSaleError as exc:
                db.rollback()
                raise HTTPException(status_code=409, detail=str(exc)) from exc
            db.refresh(sale)
            return sale

    @app.get("/api/v1/projects/{project_id}/sales", response_model=list[ProjectSaleRead])
    def list_sales(project_id: str):
        with session_scope() as db:
            if db.get(ProjectRow, project_id) is None:
                raise HTTPException(status_code=404, detail="project_not_found")
            return db.scalars(
                select(ProjectSaleRow).where(ProjectSaleRow.project_id == project_id).order_by(ProjectSaleRow.created_at)
            ).all()

    @app.post("/api/v1/projects/{project_id}/units/{unit_number}/lifecycle-events", status_code=status.HTTP_201_CREATED)
    def create_lifecycle_event(project_id: str, unit_number: str, payload: LifecycleEventCreate):
        with session_scope() as db:
            if db.get(ProjectRow, project_id) is None:
                raise HTTPException(status_code=404, detail="project_not_found")
            try:
                event = pricing_service.record_lifecycle_event(db, project_id, unit_number, payload.model_dump())
                db.commit()
            except ValueError as exc:
                db.rollback()
                raise HTTPException(status_code=422, detail=str(exc)) from exc
            except pricing_service.UnitAlreadySoldError as exc:
                db.rollback()
                raise HTTPException(status_code=409, detail=str(exc)) from exc
            return {
                "id": event.id,
                "project_id": event.project_id,
                "unit_number": event.unit_number,
                "status": event.status,
                "effective_at": event.effective_at.isoformat(),
                "source": event.source,
                "reason": event.reason,
            }

    @app.get("/api/v1/projects/{project_id}/commercial-state")
    def commercial_state(project_id: str):
        """Live commercial state -- for display/building a new snapshot only. Never
        used by any pricing or evidence path, which are always session-snapshot-scoped."""
        with session_scope() as db:
            if db.get(ProjectRow, project_id) is None:
                raise HTTPException(status_code=404, detail="project_not_found")
            latest_inventory = db.scalar(
                select(InventoryVersionRow)
                .where(InventoryVersionRow.project_id == project_id)
                .order_by(InventoryVersionRow.version_number.desc())
            )
            if latest_inventory is None:
                raise HTTPException(status_code=404, detail="inventory_version_not_found")
            unit_numbers = [u.unit_number for u in pricing_service.load_units(db, latest_inventory.id)]
            state = pricing_service.current_commercial_state(db, project_id, unit_numbers)
            return {
                "project_id": project_id,
                "inventory_version_id": latest_inventory.id,
                "units": [
                    {
                        "unit_number": unit_number,
                        "status": info["status"],
                        "effective_at": info["effective_at"].isoformat() if info["effective_at"] else None,
                        "contract_price_ils": info["sale"].contract_price_ils if info["sale"] else None,
                    }
                    for unit_number, info in sorted(state.items(), key=lambda kv: _unit_sort_key(kv[0]))
                ],
            }

    @app.post("/api/v1/projects/{project_id}/state-snapshots", status_code=status.HTTP_201_CREATED)
    def create_state_snapshot(project_id: str, payload: StateSnapshotCreate):
        with session_scope() as db:
            if db.get(ProjectRow, project_id) is None:
                raise HTTPException(status_code=404, detail="project_not_found")
            if db.get(InventoryVersionRow, payload.inventory_version_id) is None:
                raise HTTPException(status_code=404, detail="inventory_version_not_found")
            snapshot = pricing_service.create_state_snapshot(
                db, project_id, payload.inventory_version_id, as_of=payload.as_of
            )
            db.commit()
            return _state_snapshot_read(db, snapshot)

    @app.get("/api/v1/project-state-snapshots/{snapshot_id}")
    def get_state_snapshot(snapshot_id: str):
        with session_scope() as db:
            snapshot = db.get(ProjectStateSnapshotRow, snapshot_id)
            if snapshot is None:
                raise HTTPException(status_code=404, detail="project_state_snapshot_not_found")
            return _state_snapshot_read(db, snapshot)

    return app


def _import_matrix(db, project_id: str, source_filename: str, matrix: list[list[Any]]) -> InventoryVersionRead:
    normalized = normalize_inventory_rows(matrix)
    if not normalized:
        raise HTTPException(status_code=422, detail="inventory_contains_no_units")

    canonical_source = json.dumps(matrix, ensure_ascii=False, separators=(",", ":"), default=str)
    source_hash = hashlib.sha256(canonical_source.encode("utf-8")).hexdigest()

    if db.get(ProjectRow, project_id) is None:
        raise HTTPException(status_code=404, detail="project_not_found")

    existing = db.scalar(
        select(InventoryVersionRow).where(
            InventoryVersionRow.project_id == project_id,
            InventoryVersionRow.source_hash == source_hash,
        )
    )
    if existing is not None:
        return _inventory_read(existing, reused_existing=True)

    next_version = (db.scalar(
        select(func.max(InventoryVersionRow.version_number)).where(InventoryVersionRow.project_id == project_id)
    ) or 0) + 1

    version = InventoryVersionRow(
        project_id=project_id,
        version_number=next_version,
        source_filename=source_filename,
        source_hash=source_hash,
        status="complete",
    )
    db.add(version)
    db.flush()

    for record in normalized:
        unit = record.unit
        db.add(UnitRow(
            inventory_version_id=version.id,
            unit_number=unit.unit_number,
            floor=None if unit.floor is None else str(unit.floor),
            rooms=unit.rooms,
            internal_area=unit.internal_area,
            balcony_area=unit.balcony_area,
            orientation=unit.orientation,
            parking=unit.parking,
            storage=None if unit.storage is None else int(unit.storage),
            unit_type=unit.unit_type,
            notes=unit.notes,
            source_row_numbers_json=json.dumps(record.source_row_numbers),
            derivation_reasons_json=json.dumps(record.derivation_reasons, ensure_ascii=False),
            raw_rows_json=json.dumps(record.raw_rows, ensure_ascii=False, default=str),
        ))

    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="inventory_version_conflict") from exc
    db.refresh(version)
    return _inventory_read(version, reused_existing=False)


def _inventory_read(row: InventoryVersionRow, reused_existing: bool) -> InventoryVersionRead:
    # Relationship is loaded while the request session is alive.
    units = []
    for u in sorted(row.units, key=lambda x: _unit_sort_key(x.unit_number)):
        units.append(UnitRead(
            unit_number=u.unit_number,
            floor=u.floor,
            rooms=u.rooms,
            internal_area=u.internal_area,
            balcony_area=u.balcony_area,
            orientation=u.orientation,
            parking=u.parking,
            storage=None if u.storage is None else bool(u.storage),
            unit_type=u.unit_type,
            notes=u.notes,
            source_row_numbers=json.loads(u.source_row_numbers_json),
            derivation_reasons=json.loads(u.derivation_reasons_json),
        ))
    return InventoryVersionRead(
        id=row.id,
        project_id=row.project_id,
        version_number=row.version_number,
        source_filename=row.source_filename,
        source_hash=row.source_hash,
        status=row.status,
        imported_at=row.imported_at,
        reused_existing=reused_existing,
        units=units,
    )


def _snapshot_read(db, snapshot: MarketSnapshotRow) -> dict:
    runs = db.scalars(select(SourceRunRow).where(SourceRunRow.market_snapshot_id == snapshot.id)).all()
    return {
        "id": snapshot.id,
        "project_id": snapshot.project_id,
        "status": snapshot.status,
        "location": {
            "city": snapshot.location_city,
            "neighborhood": snapshot.location_neighborhood,
            "latitude": snapshot.location_latitude,
            "longitude": snapshot.location_longitude,
            "source": snapshot.location_source,
        },
        "snapshot_created_at": snapshot.snapshot_created_at.isoformat(),
        "pricing_as_of": snapshot.pricing_as_of,
        "pricing_as_of_basis": snapshot.pricing_as_of_basis,
        "pricing_as_of_excluded_sources": json.loads(snapshot.pricing_as_of_excluded_sources_json),
        "demo_disclaimer": snapshot.demo_disclaimer,
        "source_runs": [
            {
                "source_key": r.source_key,
                "lane": r.lane,
                "status": r.status,
                "raw_count": r.raw_count,
                "usable_count": r.usable_count,
                "rejected_count": r.rejected_count,
                "source_filename": r.source_filename,
                "relative_source_path": r.relative_source_path,
                "sha256": r.sha256,
                "collected_at": r.collected_at.isoformat() if r.collected_at else None,
                "collected_at_basis": r.collected_at_basis,
                "error_code": r.error_code,
                "safe_error_message": r.safe_error_message,
            }
            for r in runs
        ],
    }


def _state_snapshot_read(db, snapshot: ProjectStateSnapshotRow) -> dict:
    rows = db.scalars(
        select(ProjectStateUnitRow).where(ProjectStateUnitRow.project_state_snapshot_id == snapshot.id)
    ).all()
    units = []
    for r in sorted(rows, key=lambda x: _unit_sort_key(x.unit_number)):
        sale = db.get(ProjectSaleRow, r.project_sale_id) if r.project_sale_id else None
        units.append({
            "unit_number": r.unit_number,
            "status": r.status,
            "effective_at": r.effective_at.isoformat() if r.effective_at else None,
            "contract_price_ils": sale.contract_price_ils if sale else None,
            "contract_date": sale.contract_date if sale else None,
        })
    return {
        "id": snapshot.id,
        "project_id": snapshot.project_id,
        "inventory_version_id": snapshot.inventory_version_id,
        "as_of": snapshot.as_of.isoformat(),
        "created_at": snapshot.created_at.isoformat(),
        "units": units,
    }


def _unit_sort_key(value: str) -> tuple[int, str]:
    try:
        return (0, f"{int(value):09d}")
    except (TypeError, ValueError):
        return (1, str(value))


app = create_app()
