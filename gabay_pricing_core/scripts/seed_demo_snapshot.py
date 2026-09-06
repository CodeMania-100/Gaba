from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app_api.database import Base, build_database  # noqa: E402
from app_api.db_models import ProjectRow  # noqa: E402
from app_api.demo_snapshot import RequiredSourceMissingError, build_demo_market_snapshot  # noqa: E402
from app_api.main import DEMO_PROJECT_KEY, _import_matrix  # noqa: E402
from app_api.migrations import apply_schema_migrations  # noqa: E402
from app_api.source_manifest import default_demo_data_dir  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_XLSX = REPO_ROOT / "חוברת1.xlsx"


def main() -> None:
    parser = argparse.ArgumentParser(description="Prime the demo project with the real supplied inventory + a frozen market snapshot.")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--demo-data-dir", type=Path, default=None)
    args = parser.parse_args()

    database_url = os.getenv("DATABASE_URL", "sqlite:///./gabay_pricing.db")
    demo_data_dir = args.demo_data_dir or default_demo_data_dir()

    if not args.xlsx.exists():
        raise SystemExit(f"Supplied inventory workbook not found: {args.xlsx}")

    engine, SessionLocal = build_database(database_url)
    Base.metadata.create_all(engine)
    apply_schema_migrations(engine)
    db = SessionLocal()

    project = db.query(ProjectRow).filter(ProjectRow.demo_key == DEMO_PROJECT_KEY).one_or_none()
    if project is None:
        project = ProjectRow(
            demo_key=DEMO_PROJECT_KEY,
            name="Assignment Demo Project",
            city="אשקלון",
            neighborhood="עיר היין",
            location_source="candidate_demo_neighborhood_centroid_from_current_listings",
        )
        db.add(project)
        db.commit()
        db.refresh(project)

    workbook = openpyxl.load_workbook(args.xlsx, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    matrix = [list(row) for row in worksheet.iter_rows(values_only=True)]
    inventory = _import_matrix(db, project.id, args.xlsx.name, matrix)

    try:
        snapshot = build_demo_market_snapshot(db, project, demo_data_dir)
        db.commit()
    except RequiredSourceMissingError as exc:
        db.rollback()
        raise SystemExit(f"Required source missing: {exc}") from exc

    print(json.dumps({
        "project_id": project.id,
        "inventory_version_id": inventory.id,
        "unit_count": len(inventory.units),
        "market_snapshot_id": snapshot.id,
        "pricing_as_of": snapshot.pricing_as_of,
        "pricing_as_of_basis": snapshot.pricing_as_of_basis,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
